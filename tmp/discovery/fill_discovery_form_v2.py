from __future__ import annotations

from docx.shared import Pt

import fill_discovery_form as base


def annotate_questions(doc, questions, helper_style):
    source_titles = base.extract_original_question_titles(doc)
    annotated = []

    for entry in questions:
        number = entry["number"]
        pattern = base.exact_number_pattern(number)
        target = None
        for table in doc.tables:
            table_text = base.clean(" ".join(cell.text for row in table.rows for cell in row.cells))
            if not pattern.search(table_text):
                continue
            if not any(label in table_text for label in ("WAJIB", "DISARANKAN", "OPSIONAL", "JIKA RELEVAN")):
                continue
            for row in table.rows:
                for cell in row.cells:
                    if pattern.search(base.clean(cell.text)):
                        target = cell
                        break
                if target is not None:
                    break
            if target is not None:
                break

        if target is None:
            raise ValueError(f"Could not locate question banner {number}")

        status = target.add_paragraph(style=helper_style)
        status.paragraph_format.space_before = Pt(4)
        status.paragraph_format.space_after = Pt(2)
        run = status.add_run("Status audit: ")
        run.bold = True
        status.add_run(entry["status"])

        summary = target.add_paragraph(style=helper_style)
        summary.paragraph_format.space_after = Pt(2)
        run = summary.add_run("Jawaban berdasarkan project: ")
        run.bold = True
        summary.add_run(entry.get("short", entry["answer"]))

        pointer = target.add_paragraph(style=helper_style)
        pointer.paragraph_format.space_after = Pt(1)
        run = pointer.add_run("Sumber dan rincian: ")
        run.bold = True
        pointer.add_run(f"Lampiran Analisis — {number}.")
        annotated.append(number)

    if annotated != base.EXPECTED:
        raise ValueError(f"Annotation sequence mismatch: {annotated}")
    return source_titles


base.annotate_questions = annotate_questions

if __name__ == "__main__":
    base.main()
