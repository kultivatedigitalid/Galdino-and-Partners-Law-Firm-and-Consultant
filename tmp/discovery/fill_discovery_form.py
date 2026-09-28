from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt, RGBColor


EXPECTED = [f"{section}.{question}" for section, count in ((1, 10), (2, 9), (3, 6), (4, 8), (5, 11), (6, 6)) for question in range(1, count + 1)]

SECTION_TITLES = {
    "1": "Bagian 01 — Bisnis",
    "2": "Bagian 02 — Pelanggan",
    "3": "Bagian 03 — Positioning",
    "4": "Bagian 04 — Brand",
    "5": "Bagian 05 — Website",
    "6": "Bagian 06 — Aset",
}


def load_primary(path: Path) -> dict:
    text = path.read_text(encoding="utf-8").rstrip()
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        # The working file was deliberately produced in chunks. Complete a final
        # trailing array/object pair in memory without mutating the evidence file.
        text = text.rstrip().rstrip(",") + "\n  ]\n}\n"
        return json.loads(text)


def load_answers(primary: Path, chunks: list[Path]) -> tuple[str, list[dict], list[dict]]:
    root = load_primary(primary)
    questions = list(root["questions"])
    for chunk in chunks:
        questions.extend(json.loads(chunk.read_text(encoding="utf-8")))

    numbers = [entry["number"] for entry in questions]
    if numbers != EXPECTED:
        raise ValueError(f"Question sequence mismatch. Expected {EXPECTED}, received {numbers}")

    allowed = {"Terkonfirmasi", "Sebagian terkonfirmasi", "Belum dapat dikonfirmasi", "Pertanyaan kurang jelas"}
    for entry in [*root["identity"], *questions]:
        if entry["status"] not in allowed:
            raise ValueError(f"Unsupported status: {entry['status']}")
        for required in ("answer", "sources"):
            if not entry.get(required):
                raise ValueError(f"Missing {required} in {entry.get('number', entry.get('label'))}")

    return root["scope_note"], root["identity"], questions


def clean(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def exact_number_pattern(number: str) -> re.Pattern[str]:
    return re.compile(rf"(?<![\d.]){re.escape(number)}(?![\d.])")


def style_name(doc: Document, preferred: str, fallback: str = "Normal") -> str:
    names = {style.name for style in doc.styles}
    return preferred if preferred in names else fallback


def add_labelled_paragraph(container, label: str, body: str, style: str = "Normal", *, keep_with_next: bool = False):
    paragraph = container.add_paragraph(style=style)
    paragraph.paragraph_format.keep_with_next = keep_with_next
    label_run = paragraph.add_run(label)
    label_run.bold = True
    paragraph.add_run(body)
    return paragraph


def replace_cell(cell, lines: list[tuple[str, str]], paragraph_style: str):
    cell.text = ""
    first = cell.paragraphs[0]
    first.style = paragraph_style
    for index, (label, body) in enumerate(lines):
        paragraph = first if index == 0 else cell.add_paragraph(style=paragraph_style)
        run = paragraph.add_run(label)
        run.bold = True
        paragraph.add_run(body)


def extract_original_question_titles(doc: Document) -> dict[str, str]:
    titles: dict[str, str] = {}
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = clean(cell.text)
                match = re.search(r"(?<![\d.])([1-6]\.(?:10|11|[1-9]))(?![\d.])\s+(.+)", text)
                if match and match.group(1) not in titles:
                    titles[match.group(1)] = match.group(2).strip()
    return titles


def annotate_questions(doc: Document, questions: list[dict], helper_style: str) -> dict[str, str]:
    source_titles = extract_original_question_titles(doc)
    annotated: list[str] = []

    for entry in questions:
        number = entry["number"]
        pattern = exact_number_pattern(number)
        target = None
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    text = clean(cell.text)
                    if pattern.search(text) and re.search(r"WAJIB|DISARANKAN|OPSIONAL|JIKA RELEVAN", text):
                        target = cell
                        break
                if target is not None:
                    break
            if target is not None:
                break
        if target is None:
            raise ValueError(f"Could not locate question banner {number}")

        status = target.add_paragraph(style=helper_style)
        status.paragraph_format.space_before = Pt(4)
        status.paragraph_format.space_after = Pt(2)
        run = status.add_run("Status audit: ")
        run.bold = True
        status.add_run(entry["status"])

        summary = target.add_paragraph(style=helper_style)
        summary.paragraph_format.space_after = Pt(2)
        run = summary.add_run("Jawaban berdasarkan project: ")
        run.bold = True
        summary.add_run(entry.get("short", entry["answer"]))

        pointer = target.add_paragraph(style=helper_style)
        pointer.paragraph_format.space_after = Pt(1)
        run = pointer.add_run("Sumber dan rincian: ")
        run.bold = True
        pointer.add_run(f"Lampiran Analisis — {number}.")
        annotated.append(number)

    if annotated != EXPECTED:
        raise ValueError(f"Annotation sequence mismatch: {annotated}")
    return source_titles


def fill_identity(doc: Document, identity: list[dict], helper_style: str):
    table = None
    for candidate in doc.tables:
        if len(candidate.rows) == 7 and "Nama perusahaan / brand" in candidate.cell(0, 0).text:
            table = candidate
            break
    if table is None:
        raise ValueError("Identity table not found")
    if len(identity) != 7:
        raise ValueError("Identity answer count must be 7")

    for row, entry in zip(table.rows, identity):
        replace_cell(
            row.cells[-1],
            [
                ("Status: ", entry["status"]),
                ("Jawaban berdasarkan project: ", entry["answer"]),
                ("Sumber dan rincian: ", f"Lampiran Analisis — Identitas Proyek / {entry['label']}.")
            ],
            helper_style,
        )


def annotate_attestations(doc: Document, helper_style: str):
    final_table = None
    for table in doc.tables:
        if "Jawaban menggambarkan kondisi perusahaan saat ini" in clean(table.cell(0, 0).text):
            final_table = table
            break
    if final_table is None:
        raise ValueError("Final confirmation table not found")

    notes = [
        "Belum dapat dikonfirmasi dari codebase; kondisi bisnis memerlukan validasi stakeholder.",
        "Belum dapat dikonfirmasi; link folder aset eksternal tidak ditemukan.",
        "Sebagian terkonfirmasi; placeholder dan batas publikasi dicatat, tetapi attestation tetap memerlukan persetujuan stakeholder.",
        "Belum dapat dikonfirmasi; PIC utama dan approver final tidak ditetapkan di project.",
    ]
    for row, note in zip(final_table.rows, notes):
        paragraph = row.cells[0].add_paragraph(style=helper_style)
        run = paragraph.add_run("Status audit: ")
        run.bold = True
        paragraph.add_run(note)

    filler_table = None
    for table in doc.tables:
        if len(table.rows) == 3 and "Nama pengisi" in clean(table.cell(0, 0).text):
            filler_table = table
            break
    if filler_table is None:
        raise ValueError("Filler identity table not found")
    for row in filler_table.rows:
        replace_cell(row.cells[-1], [("Status: ", "Belum dapat dikonfirmasi dari project.")], helper_style)

    note_table = None
    for table in doc.tables:
        if len(table.rows) == 1 and clean(table.cell(0, 0).text) == "Klik dan ketik jawaban Anda...":
            # The final matching one is the additional-notes response table.
            note_table = table
    if note_table is None:
        raise ValueError("Additional notes table not found")
    replace_cell(
        note_table.cell(0, 0),
        [
            ("Catatan audit: ", "Dokumen diisi hanya dari bukti repository. Tidak ada sumber eksternal yang digunakan."),
            ("Batasan: ", "Placeholder, mock, dummy, concept, [DATA], [VERIFY], pending, dan awaiting approval tidak diperlakukan sebagai kondisi bisnis terverifikasi."),
        ],
        helper_style,
    )


def add_sources(doc: Document, sources: list[str], normal_style: str):
    heading = doc.add_paragraph(style=normal_style)
    heading.paragraph_format.keep_with_next = True
    heading.add_run("Sumber dalam project:").bold = True
    for source in sources:
        paragraph = doc.add_paragraph(style="List Bullet" if "List Bullet" in {s.name for s in doc.styles} else normal_style)
        paragraph.add_run(source)


def add_entry(doc: Document, heading: str, entry: dict, heading_style: str, normal_style: str):
    paragraph = doc.add_paragraph(style=heading_style)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run(heading)

    add_labelled_paragraph(doc, "Status: ", entry["status"], normal_style)
    add_labelled_paragraph(doc, "Jawaban berdasarkan project: ", entry["answer"], normal_style)
    add_sources(doc, entry["sources"], normal_style)

    if entry.get("missing"):
        add_labelled_paragraph(doc, "Informasi yang belum tersedia atau belum pasti: ", entry["missing"], normal_style)
    add_labelled_paragraph(doc, "Rekomendasi: ", entry.get("recommendation", "Tidak ada rekomendasi tambahan; gunakan kondisi terkonfirmasi di atas."), normal_style)
    if entry.get("additional"):
        add_labelled_paragraph(doc, "Data tambahan yang dibutuhkan: ", entry["additional"], normal_style)


def append_analysis(doc: Document, scope_note: str, identity: list[dict], questions: list[dict], source_titles: dict[str, str]):
    title_style = style_name(doc, "K Title")
    subtitle_style = style_name(doc, "K Subtitle")
    section_style = style_name(doc, "K Section")
    heading_style = style_name(doc, "K Heading")
    normal_style = style_name(doc, "Normal")

    page_break = doc.add_paragraph()
    page_break.add_run().add_break(WD_BREAK.PAGE)

    title = doc.add_paragraph(style=title_style)
    title.add_run("Lampiran Analisis Berdasarkan Project")
    subtitle = doc.add_paragraph(style=subtitle_style)
    subtitle.add_run("Jawaban lengkap, status verifikasi, sumber, gap informasi, rekomendasi, dan data tambahan yang dibutuhkan.")

    add_labelled_paragraph(doc, "Ruang lingkup audit: ", scope_note, normal_style)
    add_labelled_paragraph(
        doc,
        "Legenda status: ",
        "Terkonfirmasi = terbukti langsung; Sebagian terkonfirmasi = hanya sebagian bukti tersedia; Belum dapat dikonfirmasi = bukti tidak cukup; Pertanyaan kurang jelas = makna membutuhkan konteks tambahan.",
        normal_style,
    )
    add_labelled_paragraph(
        doc,
        "Prinsip pembacaan: ",
        "Implementasi aktual diprioritaskan atas dokumentasi. Rencana, TODO, placeholder, mock, dummy, generated concept, dan data menunggu approval disebutkan secara eksplisit dan tidak diperlakukan sebagai fitur/fakta final.",
        normal_style,
    )

    section = doc.add_paragraph(style=section_style)
    section.add_run("Identitas Proyek")
    for entry in identity:
        add_entry(doc, entry["label"], entry, heading_style, normal_style)

    current_section = None
    for entry in questions:
        section_number = entry["number"].split(".")[0]
        if section_number != current_section:
            section = doc.add_paragraph(style=section_style)
            section.paragraph_format.page_break_before = True
            section.add_run(SECTION_TITLES[section_number])
            current_section = section_number
        title = source_titles.get(entry["number"], entry["title"])
        add_entry(doc, f"{entry['number']}  {title}", entry, heading_style, normal_style)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("primary", type=Path)
    parser.add_argument("chunks", nargs="+", type=Path)
    args = parser.parse_args()

    scope_note, identity, questions = load_answers(args.primary, args.chunks)
    doc = Document(args.source)
    original_tables = len(doc.tables)
    original_sections = len(doc.sections)

    helper_style = style_name(doc, "K Helper")
    fill_identity(doc, identity, helper_style)
    source_titles = annotate_questions(doc, questions, helper_style)
    annotate_attestations(doc, helper_style)
    append_analysis(doc, scope_note, identity, questions, source_titles)

    if len(doc.tables) != original_tables:
        raise ValueError("Original table count changed")
    if len(doc.sections) != original_sections:
        raise ValueError("Original section count changed")

    doc.core_properties.title = "Galdino & Partner FASE 1 — Discovery — Terisi Berdasarkan Project"
    doc.core_properties.subject = "Project discovery audit based only on repository evidence"
    doc.core_properties.comments = "Generated from project evidence; source form preserved and detailed analysis appended."
    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)
    print(json.dumps({
        "output": str(args.output.resolve()),
        "questions": len(questions),
        "identity_fields": len(identity),
        "tables_preserved": len(doc.tables),
        "sections_preserved": len(doc.sections),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
