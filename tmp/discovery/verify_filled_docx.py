from __future__ import annotations

import argparse
import hashlib
import json
import re
import zipfile
from pathlib import Path

from docx import Document


EXPECTED = [f"{section}.{question}" for section, count in ((1, 10), (2, 9), (3, 6), (4, 8), (5, 11), (6, 6)) for question in range(1, count + 1)]
EXPECTED_SOURCE_SHA256 = "2ff093b8653baaa0d21f8e92ddc6fca15b7dd1896e6518069e07592c1f51eb47"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def clean(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def section_signature(section):
    return {
        "page_width": section.page_width,
        "page_height": section.page_height,
        "left_margin": section.left_margin,
        "right_margin": section.right_margin,
        "top_margin": section.top_margin,
        "bottom_margin": section.bottom_margin,
        "header_distance": section.header_distance,
        "footer_distance": section.footer_distance,
        "orientation": int(section.orientation),
    }


def package_diff(source: Path, output: Path):
    with zipfile.ZipFile(source) as src, zipfile.ZipFile(output) as out:
        if src.testzip() is not None or out.testzip() is not None:
            raise AssertionError("DOCX zip integrity failed")
        src_names = set(src.namelist())
        out_names = set(out.namelist())
        missing = sorted(src_names - out_names)
        added = sorted(out_names - src_names)
        changed = sorted(name for name in src_names & out_names if src.read(name) != out.read(name))
    return {"missing_parts": missing, "added_parts": added, "changed_parts": changed}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("report", type=Path)
    args = parser.parse_args()

    source_sha = sha256(args.source)
    assert source_sha == EXPECTED_SOURCE_SHA256, f"Source form changed: {source_sha}"

    source = Document(args.source)
    output = Document(args.output)
    assert len(source.sections) == len(output.sections) == 1
    assert len(source.tables) == len(output.tables) == 127
    assert [len(table.rows) for table in source.tables] == [len(table.rows) for table in output.tables]
    assert section_signature(source.sections[0]) == section_signature(output.sections[0])

    banner_checks = {}
    for number in EXPECTED:
        pattern = re.compile(rf"(?<![\d.]){re.escape(number)}(?![\d.])")
        matches = []
        for table in output.tables:
            text = clean(" ".join(cell.text for row in table.rows for cell in row.cells))
            if pattern.search(text) and any(label in text for label in ("WAJIB", "DISARANKAN", "OPSIONAL", "JIKA RELEVAN")):
                matches.append(text)
        assert len(matches) == 1, f"Question banner count for {number}: {len(matches)}"
        assert matches[0].count("Status audit:") == 1
        assert "Jawaban berdasarkan project:" in matches[0]
        assert f"Lampiran Analisis — {number}." in matches[0]
        banner_checks[number] = "ok"

    appendix_started = False
    appendix_headings = []
    appendix_text_parts = []
    for paragraph in output.paragraphs:
        text = clean(paragraph.text)
        if text == "Lampiran Analisis Berdasarkan Project":
            appendix_started = True
        if appendix_started:
            appendix_text_parts.append(text)
            if paragraph.style and paragraph.style.name == "K Heading":
                match = re.match(r"^([1-6]\.(?:10|11|[1-9]))\s+", text)
                if match:
                    appendix_headings.append(match.group(1))
    appendix_text = "\n".join(appendix_text_parts)
    assert appendix_started
    assert appendix_headings == EXPECTED, appendix_headings
    assert appendix_text.count("Status:") >= 57
    assert appendix_text.count("Jawaban berdasarkan project:") >= 57
    assert appendix_text.count("Sumber dalam project:") >= 57
    assert appendix_text.count("Rekomendasi:") >= 57
    assert "Ruang lingkup audit:" in appendix_text
    assert "Implementasi aktual diprioritaskan atas dokumentasi" in appendix_text

    identity_table = next(table for table in output.tables if len(table.rows) == 7 and "Nama perusahaan / brand" in table.cell(0, 0).text)
    for row in identity_table.rows:
        answer_text = clean(row.cells[-1].text)
        assert "Status:" in answer_text
        assert "Jawaban berdasarkan project:" in answer_text
        assert "Sumber dan rincian:" in answer_text

    final_table = next(table for table in output.tables if "Jawaban menggambarkan kondisi perusahaan saat ini" in table.cell(0, 0).text)
    assert len(final_table.rows) == 4
    assert all("Status audit:" in clean(row.cells[0].text) for row in final_table.rows)

    filler = next(table for table in output.tables if len(table.rows) == 3 and "Nama pengisi" in table.cell(0, 0).text)
    assert all("Belum dapat dikonfirmasi dari project" in clean(row.cells[-1].text) for row in filler.rows)

    document_text = "\n".join(clean(paragraph.text) for paragraph in output.paragraphs)
    document_text += "\n" + "\n".join(clean(cell.text) for table in output.tables for row in table.rows for cell in row.cells)
    assert ":codex-file-citation" not in document_text
    for env_name in (
        "PUBLIC_SITE_URL", "RESEND_API_KEY", "CONTACT_TO_EMAIL", "CONTACT_FROM_EMAIL",
        "UPSTASH_REDIS_REST_URL", "UPSTASH_REDIS_REST_TOKEN", "CONTACT_FORM_DRY_RUN",
    ):
        assert f"{env_name}=" not in document_text

    pkg = package_diff(args.source, args.output)
    assert not pkg["missing_parts"]
    assert not pkg["added_parts"]
    allowed_changes = {"word/document.xml", "docProps/core.xml"}
    unexpected_changes = sorted(set(pkg["changed_parts"]) - allowed_changes)
    assert not unexpected_changes, unexpected_changes

    report = {
        "source_sha256": source_sha,
        "output_sha256": sha256(args.output),
        "output_size_bytes": args.output.stat().st_size,
        "sections": len(output.sections),
        "tables": len(output.tables),
        "question_banners_annotated": len(banner_checks),
        "appendix_question_headings": len(appendix_headings),
        "identity_fields_filled": len(identity_table.rows),
        "final_attestations_annotated": len(final_table.rows),
        "package_diff": pkg,
        "result": "PASS",
    }
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
