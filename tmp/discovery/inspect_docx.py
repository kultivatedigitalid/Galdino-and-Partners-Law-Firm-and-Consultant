from __future__ import annotations

import hashlib
import json
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


def iter_blocks(parent):
    parent_el = parent.element.body if hasattr(parent, "element") else parent._tc
    for child in parent_el.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent)


def paragraph_record(paragraph: Paragraph) -> dict:
    return {
        "type": "paragraph",
        "text": paragraph.text,
        "style": paragraph.style.name if paragraph.style else None,
        "runs": [
            {
                "text": run.text,
                "bold": run.bold,
                "italic": run.italic,
                "underline": bool(run.underline),
                "font": run.font.name,
                "size_pt": run.font.size.pt if run.font.size else None,
            }
            for run in paragraph.runs
        ],
    }


def table_record(table: Table) -> dict:
    return {
        "type": "table",
        "style": table.style.name if table.style else None,
        "rows": [
            [
                [paragraph_record(p) for p in cell.paragraphs]
                for cell in row.cells
            ]
            for row in table.rows
        ],
    }


def part_paragraphs(part) -> list[dict]:
    return [paragraph_record(p) for p in part.paragraphs]


def main() -> None:
    source = Path(sys.argv[1]).resolve()
    out_dir = Path(sys.argv[2]).resolve()
    out_dir.mkdir(parents=True, exist_ok=True)

    doc = Document(source)
    blocks = []
    for block in iter_blocks(doc):
        blocks.append(paragraph_record(block) if isinstance(block, Paragraph) else table_record(block))

    sections = []
    for idx, section in enumerate(doc.sections, start=1):
        sections.append(
            {
                "index": idx,
                "page_width_inches": section.page_width.inches,
                "page_height_inches": section.page_height.inches,
                "top_margin_inches": section.top_margin.inches,
                "bottom_margin_inches": section.bottom_margin.inches,
                "left_margin_inches": section.left_margin.inches,
                "right_margin_inches": section.right_margin.inches,
                "header_distance_inches": section.header_distance.inches,
                "footer_distance_inches": section.footer_distance.inches,
                "different_first_page": section.different_first_page_header_footer,
                "header": part_paragraphs(section.header),
                "footer": part_paragraphs(section.footer),
                "first_page_header": part_paragraphs(section.first_page_header),
                "first_page_footer": part_paragraphs(section.first_page_footer),
            }
        )

    sha256 = hashlib.sha256(source.read_bytes()).hexdigest()
    with zipfile.ZipFile(source) as package:
        parts = [
            {"path": info.filename, "size": info.file_size, "crc": info.CRC}
            for info in package.infolist()
        ]
        document_xml = package.read("word/document.xml").decode("utf-8", errors="replace")

    payload = {
        "source": str(source),
        "sha256": sha256,
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables),
        "inline_shape_count": len(doc.inline_shapes),
        "sections": sections,
        "blocks": blocks,
        "package_parts": parts,
        "document_xml_contains_textbox": "txbxContent" in document_xml,
    }
    (out_dir / "form_structure.json").write_text(
        json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8"
    )

    lines = []
    for idx, block in enumerate(blocks, start=1):
        if block["type"] == "paragraph":
            lines.append(f"P{idx:03d} [{block['style']}] {block['text']}")
        else:
            lines.append(f"T{idx:03d} [{block['style']}] rows={len(block['rows'])}")
            for row_idx, row in enumerate(block["rows"], start=1):
                cells = []
                for cell in row:
                    cells.append(" / ".join(p["text"] for p in cell if p["text"]))
                lines.append(f"  R{row_idx:02d}: " + " || ".join(cells))
    (out_dir / "form_text.txt").write_text("\n".join(lines), encoding="utf-8")


if __name__ == "__main__":
    main()
