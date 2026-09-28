from __future__ import annotations

from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


PATH = Path(r"C:\Users\Joshua\Downloads\GP_Fase2_Website_Content_Sheet_Final_Aligned_2026-08-06.xlsx")
wb = load_workbook(PATH, data_only=False, read_only=True)

id_ws = wb["03_Content_ID"]
en_ws = wb["04_Content_EN"]


def records(ws):
    result = []
    for row in range(5, 49):
        values = [ws.cell(row, col).value for col in range(1, 18)]
        if values[0] not in (None, ""):
            result.append(values)
    return result


id_records = records(id_ws)
en_records = records(en_ws)
id_ids = [row[0] for row in id_records]
en_ids = [row[0] for row in en_records]

print("ID records:", len(id_records))
print("EN records:", len(en_records))
print("ID unique:", len(set(id_ids)))
print("EN unique:", len(set(en_ids)))
print("Same ordered IDs:", id_ids == en_ids)
print("Missing in EN:", sorted(set(id_ids) - set(en_ids)))
print("Missing in ID:", sorted(set(en_ids) - set(id_ids)))
print("Duplicate ID IDs:", [key for key, count in Counter(id_ids).items() if count > 1])
print("Duplicate EN IDs:", [key for key, count in Counter(en_ids).items() if count > 1])

for label, rows in (("ID", id_records), ("EN", en_records)):
    missing_by_column = []
    for column in range(17):
        missing = sum(1 for row in rows if row[column] in (None, ""))
        missing_by_column.append(missing)
    print(label, "missing cells by A:Q:", missing_by_column)
