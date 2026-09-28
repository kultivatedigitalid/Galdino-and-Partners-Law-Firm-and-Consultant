from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook


PATH = Path(r"C:\Users\Joshua\Downloads\GP_Fase2_Website_Content_Sheet_Final_Aligned_2026-08-06.xlsx")
wb = load_workbook(PATH, data_only=False, read_only=False)

keywords = (
    "website content sheet",
    "global",
    "headline",
    "body",
    "supporting",
    "proof",
    "bukti",
    "asset id",
    "reference id",
    "section family",
    "variant",
    "structure note",
    "ready for fact check",
    "fact checked",
    "copy approved",
    "mapped",
    "fact check",
    "approval",
    "review",
    "source",
    "sumber",
)

result = {"file": str(PATH), "sheets": []}
for ws in wb.worksheets:
    nonempty_rows = []
    keyword_hits = []
    status_counts: dict[str, int] = {}
    for row in ws.iter_rows():
        values = [cell.value for cell in row]
        if any(value not in (None, "") for value in values):
            nonempty_rows.append(row[0].row)
        for cell in row:
            if cell.value in (None, ""):
                continue
            text = str(cell.value)
            lower = text.casefold()
            if any(keyword in lower for keyword in keywords):
                keyword_hits.append({"cell": cell.coordinate, "value": text[:800]})
            if any(token in lower for token in ("approved", "checked", "mapped", "draft", "review", "missing", "partial", "ready")):
                status_counts[text] = status_counts.get(text, 0) + 1
    result["sheets"].append({
        "name": ws.title,
        "dimensions": ws.calculate_dimension(),
        "nonempty_row_count": len(nonempty_rows),
        "first_nonempty_row": min(nonempty_rows) if nonempty_rows else None,
        "last_nonempty_row": max(nonempty_rows) if nonempty_rows else None,
        "merged_ranges": [str(rng) for rng in ws.merged_cells.ranges],
        "keyword_hits": keyword_hits,
        "status_counts": status_counts,
    })

print(json.dumps(result, ensure_ascii=False, indent=2, default=str))
