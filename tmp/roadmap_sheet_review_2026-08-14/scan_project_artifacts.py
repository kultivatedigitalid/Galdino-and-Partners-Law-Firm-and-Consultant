from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from openpyxl import load_workbook


ROOT = Path(r"C:\Users\Joshua\OneDrive\Documents\Law\outputs\galdino_partner")
KEYWORDS = (
    "website content sheet",
    "copy approved",
    "ready for fact check",
    "fact checked",
    "asset id",
    "reference id",
    "section family",
    "structure note",
    "mapped",
    "headline",
    "supporting points",
)


def matched(value: object) -> bool:
    text = str(value).casefold()
    return any(keyword in text for keyword in KEYWORDS)


result = {"xlsx": [], "docx": []}

for path in sorted(ROOT.rglob("*.xlsx")):
    try:
        wb = load_workbook(path, data_only=False, read_only=True)
        hits = []
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value not in (None, "") and matched(cell.value):
                        hits.append({"sheet": ws.title, "cell": cell.coordinate, "value": str(cell.value)[:500]})
        result["xlsx"].append({"file": str(path.relative_to(ROOT)), "sheets": wb.sheetnames, "hits": hits})
    except Exception as exc:
        result["xlsx"].append({"file": str(path.relative_to(ROOT)), "error": str(exc)})

for path in sorted(ROOT.rglob("*.docx")):
    try:
        doc = Document(path)
        hits = []
        for index, paragraph in enumerate(doc.paragraphs, start=1):
            if paragraph.text and matched(paragraph.text):
                hits.append({"location": f"paragraph {index}", "value": paragraph.text[:500]})
        for table_index, table in enumerate(doc.tables, start=1):
            for row_index, row in enumerate(table.rows, start=1):
                for cell_index, cell in enumerate(row.cells, start=1):
                    if cell.text and matched(cell.text):
                        hits.append({"location": f"table {table_index} row {row_index} cell {cell_index}", "value": cell.text[:500]})
        result["docx"].append({"file": str(path.relative_to(ROOT)), "hits": hits})
    except Exception as exc:
        result["docx"].append({"file": str(path.relative_to(ROOT)), "error": str(exc)})

print(json.dumps(result, ensure_ascii=False, indent=2))
