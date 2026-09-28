from __future__ import annotations

import json
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


PATH = Path(r"C:\Users\Joshua\Downloads\GP_Fase2_Website_Content_Sheet_Final_Aligned_2026-08-06.xlsx")
wb = load_workbook(PATH, data_only=False, read_only=True)


def row_values(ws, row_number: int, max_col: int) -> list[object]:
    return [ws.cell(row_number, col).value for col in range(1, max_col + 1)]


result: dict[str, object] = {"sheets": wb.sheetnames, "ranges": {}}

for sheet_name, rows, max_col in (
    ("00_Dashboard", list(range(1, 13)) + list(range(20, 26)) + list(range(47, 55)), 10),
    ("01_GLOBAL", list(range(1, 11)) + list(range(20, 47)), 10),
    ("02_Page_Questions", list(range(1, 21)), 10),
    ("03_Content_ID", list(range(1, 9)), 17),
    ("04_Content_EN", list(range(1, 9)), 17),
    ("05_Proof_Placeholders", list(range(1, 27)), 10),
    ("06_Positioning_Tone_QA", list(range(1, 18)), 12),
    ("07_Approval_Actions", list(range(1, 23)), 9),
    ("08_Source_Notes", list(range(1, 17)), 5),
):
    ws = wb[sheet_name]
    result["ranges"][sheet_name] = [{"row": row, "values": row_values(ws, row, max_col)} for row in rows]

for sheet_name in ("03_Content_ID", "04_Content_EN"):
    ws = wb[sheet_name]
    headers = row_values(ws, 4, 17)
    columns: dict[str, object] = {}
    for col_index, header in enumerate(headers, start=1):
        values = [ws.cell(row, col_index).value for row in range(5, 49)]
        nonempty = [value for value in values if value not in (None, "")]
        columns[str(header)] = {
            "column": ws.cell(4, col_index).column_letter,
            "nonempty": len(nonempty),
            "distinct_counts": Counter(str(value) for value in nonempty).most_common(30),
        }
    result.setdefault("content_columns", {})[sheet_name] = columns

print(json.dumps(result, ensure_ascii=False, indent=2, default=str))
