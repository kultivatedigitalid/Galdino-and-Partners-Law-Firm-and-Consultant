from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook


PATH = Path(r"C:\Users\Joshua\Downloads\GP_Fase2_Website_Content_Sheet_Final_Aligned_2026-08-06.xlsx")
wb = load_workbook(PATH, data_only=False, read_only=True)


def print_rows(sheet_name: str, start: int, end: int, max_col: int) -> None:
    ws = wb[sheet_name]
    print(f"\n## {sheet_name} rows {start}:{end}")
    for row_number in range(start, end + 1):
        values = [ws.cell(row_number, column).value for column in range(1, max_col + 1)]
        print(f"{row_number}: " + " | ".join("" if value is None else str(value).replace("\n", " / ") for value in values))


print_rows("03_Content_ID", 1, 6, 17)
print_rows("04_Content_EN", 1, 6, 17)
print_rows("06_Positioning_Tone_QA", 1, 17, 12)
print_rows("07_Approval_Actions", 1, 22, 9)
