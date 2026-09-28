from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(r"C:\Users\Joshua\OneDrive\Documents\Law\outputs\galdino_partner\alignment_2026-08-05")
TARGETS = {
    "GP_Fase2_Architecture_Workbook_Aligned_2026-08-05.xlsx": [
        "Act 6 Page Questions",
        "Job Question Alignment",
        "Content Dependencies",
    ],
    "GP_Alignment_QA_Checklist_2026-08-05.xlsx": [
        "Website Baseline",
        "Delivery Checklist",
    ],
}

result = {}
for filename, sheet_names in TARGETS.items():
    wb = load_workbook(ROOT / filename, data_only=False, read_only=True)
    file_result = {}
    for sheet_name in sheet_names:
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows():
            values = [cell.value for cell in row]
            if any(value not in (None, "") for value in values):
                rows.append({"row": row[0].row, "values": values})
        file_result[sheet_name] = {"max_row": ws.max_row, "max_column": ws.max_column, "rows": rows}
    result[filename] = file_result

print(json.dumps(result, ensure_ascii=False, indent=2, default=str))
