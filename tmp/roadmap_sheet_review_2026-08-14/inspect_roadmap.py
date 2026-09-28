from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook


PATH = Path(r"C:\Users\Joshua\OneDrive\Documents\Law\tmp\roadmap_sheet_review_2026-08-14\Roadmap website Galdino & Partner - Kultivate.xlsx")
wb = load_workbook(PATH, data_only=False, read_only=False)

result = {"file": str(PATH), "sheets": []}
for ws in wb.worksheets:
    cells = []
    for row in ws.iter_rows():
        for cell in row:
            if cell.value not in (None, ""):
                cells.append({
                    "cell": cell.coordinate,
                    "value": cell.value,
                    "data_type": cell.data_type,
                    "number_format": cell.number_format,
                })
    validations = []
    if ws.data_validations:
        for dv in ws.data_validations.dataValidation:
            validations.append({
                "type": dv.type,
                "formula1": dv.formula1,
                "formula2": dv.formula2,
                "ranges": str(dv.sqref),
            })
    result["sheets"].append({
        "name": ws.title,
        "max_row": ws.max_row,
        "max_column": ws.max_column,
        "merged_ranges": [str(item) for item in ws.merged_cells.ranges],
        "validations": validations,
        "cells": cells,
    })

print(json.dumps(result, ensure_ascii=False, indent=2, default=str))
