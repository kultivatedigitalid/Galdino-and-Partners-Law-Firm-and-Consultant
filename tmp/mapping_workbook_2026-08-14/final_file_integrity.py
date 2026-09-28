import json
import zipfile
from pathlib import Path
from openpyxl import load_workbook

path = Path(r"C:\Users\Joshua\OneDrive\Documents\Law\outputs\galdino_partner\fase2\GP_Fase2_Website_Content_Sheet_Mapping_Aligned_2026-08-14.xlsx")
expected = [
    "00_Dashboard", "01_GLOBAL", "02_Page_Questions", "03_Content_ID", "04_Content_EN",
    "05_Proof_Placeholders", "06_Positioning_Tone_QA", "07_Approval_Actions", "08_Source_Notes",
    "09_Page_Section_Mapping", "10_Asset_Register", "11_Reference_Register", "12_Mapping_QA",
]

with zipfile.ZipFile(path, "r") as archive:
    bad_zip_member = archive.testzip()

workbook = load_workbook(path, read_only=False, data_only=False)
mapping = workbook["09_Page_Section_Mapping"]
assets = workbook["10_Asset_Register"]
references = workbook["11_Reference_Register"]
qa = workbook["12_Mapping_QA"]

formula_count = 0
for sheet in workbook.worksheets:
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                formula_count += 1

result = {
    "path": str(path),
    "sizeBytes": path.stat().st_size,
    "zipIntegrity": "PASS" if bad_zip_member is None else f"FAIL: {bad_zip_member}",
    "sheetNamesMatch": workbook.sheetnames == expected,
    "sheetCount": len(workbook.sheetnames),
    "mappingRows": mapping.max_row - 4,
    "assetRows": assets.max_row - 4,
    "referenceRows": references.max_row - 4,
    "qaRows": 13,
    "mappingHeader": mapping["A4"].value,
    "mappingLastId": mapping["C50"].value,
    "serviceDetailsFormula": qa["G14"].value,
    "finalMappedFormula": qa["B25"].value,
    "formulaCount": formula_count,
    "row3Heights": {sheet.title: sheet.row_dimensions[3].height for sheet in [mapping, assets, references, qa]},
    "mergedNotesRemoved": all(not any(str(rng).startswith("A2:") or str(rng).startswith("A3:") for rng in sheet.merged_cells.ranges) for sheet in [mapping, assets, references, qa]),
}
print(json.dumps(result, ensure_ascii=False, indent=2))

if bad_zip_member is not None or workbook.sheetnames != expected or mapping.max_row != 50 or assets.max_row != 22 or references.max_row != 31:
    raise SystemExit(1)
