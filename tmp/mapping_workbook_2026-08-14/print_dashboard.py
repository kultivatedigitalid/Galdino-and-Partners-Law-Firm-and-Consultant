from openpyxl import load_workbook

path = r"C:\Users\Joshua\OneDrive\Documents\Law\tmp\mapping_workbook_2026-08-14\source_sanitized_for_artifact_tool.xlsx"
workbook = load_workbook(path, read_only=True, data_only=False)
sheet = workbook["00_Dashboard"]
for row in range(1, 70):
    values = [sheet.cell(row, col).value for col in range(1, 11)]
    if any(value not in (None, "") for value in values):
        print(row, values)
