import json
import zipfile
from pathlib import Path
from openpyxl import load_workbook

path = Path(r"C:\Users\Joshua\OneDrive\Documents\Law\outputs\galdino_partner\fase2\GP_Fase2_Website_Content_Sheet_Mapping_Checklist_Evidence_2026-08-14.xlsx")
expected = [
    "00_Dashboard", "01_GLOBAL", "02_Page_Questions", "03_Content_ID", "04_Content_EN",
    "05_Proof_Placeholders", "06_Positioning_Tone_QA", "07_Approval_Actions", "08_Source_Notes",
    "09_Page_Section_Mapping", "10_Asset_Register", "11_Reference_Register", "12_Mapping_QA",
    "13_Checklist_Evidence",
]
with zipfile.ZipFile(path, "r") as archive:
    bad_zip_member = archive.testzip()

wb = load_workbook(path, read_only=False, data_only=False)
ws = wb["13_Checklist_Evidence"]
data = [[ws.cell(row, col).value for col in range(1, 7)] for row in range(5, 34)]
status = [row[4] for row in data]
result = {
    "path": str(path),
    "sizeBytes": path.stat().st_size,
    "zipIntegrity": "PASS" if bad_zip_member is None else f"FAIL: {bad_zip_member}",
    "sheetNamesMatch": wb.sheetnames == expected,
    "sheetCount": len(wb.sheetnames),
    "checklistRows": len(data),
    "canCheck": status.count("BISA DICENTANG SEKARANG"),
    "cannotCheck": status.count("BELUM BISA DICENTANG"),
    "boxesRemainBlank": all(row[3] == "[ ]" for row in data),
    "allExplanationsPresent": all(bool(row[5]) for row in data),
    "uniqueRefs": len({row[1] for row in data}) == 29,
}
print(json.dumps(result, ensure_ascii=False, indent=2))
checks = [
    bad_zip_member is None, wb.sheetnames == expected, len(data) == 29,
    result["canCheck"] == 17, result["cannotCheck"] == 12,
    result["boxesRemainBlank"], result["allExplanationsPresent"], result["uniqueRefs"],
]
raise SystemExit(0 if all(checks) else 1)
