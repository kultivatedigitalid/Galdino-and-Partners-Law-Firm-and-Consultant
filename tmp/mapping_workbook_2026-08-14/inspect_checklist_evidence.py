from pathlib import Path
from collections import Counter

from openpyxl import load_workbook


PATH = Path(r"C:\Users\Joshua\OneDrive\Documents\Law\outputs\galdino_partner\fase2\GP_Fase2_Website_Content_Sheet_Mapping_Aligned_2026-08-14.xlsx")
wb = load_workbook(PATH, data_only=False, read_only=True)

for sheet_name in ["03_Content_ID", "05_Proof_Placeholders", "06_Positioning_Tone_QA", "07_Approval_Actions", "09_Page_Section_Mapping", "10_Asset_Register", "11_Reference_Register", "12_Mapping_QA"]:
    ws = wb[sheet_name]
    headers = [ws.cell(4, col).value for col in range(1, ws.max_column + 1)]
    print(f"\n[{sheet_name}] rows={ws.max_row} cols={ws.max_column}")
    print("HEADERS:", headers)
    if sheet_name == "09_Page_Section_Mapping":
        for idx, header in enumerate(headers, start=1):
            values = [ws.cell(row, idx).value for row in range(5, ws.max_row + 1)]
            nonblank = sum(value not in (None, "") for value in values)
            counts = Counter(str(value) for value in values if value not in (None, ""))
            print(f"{idx:02d} {header}: nonblank={nonblank}; top={counts.most_common(5)}")

