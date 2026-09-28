from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook


PATH = Path(r"C:\Users\Joshua\Downloads\GP_Fase2_Website_Content_Sheet_Final_Aligned_2026-08-06.xlsx")
OUTPUT = Path(r"C:\Users\Joshua\OneDrive\Documents\Law\tmp\mapping_workbook_2026-08-14\content_sheet_data.json")
wb = load_workbook(PATH, data_only=False, read_only=True)


def table(sheet_name: str, header_row: int, start_row: int, end_row: int, max_col: int):
    ws = wb[sheet_name]
    headers = [ws.cell(header_row, col).value for col in range(1, max_col + 1)]
    rows = []
    for row in range(start_row, end_row + 1):
        values = [ws.cell(row, col).value for col in range(1, max_col + 1)]
        if values[0] in (None, ""):
            continue
        rows.append({str(headers[index]): values[index] for index in range(max_col)})
    return rows


data = {
    "global": table("01_GLOBAL", 4, 5, 45, 10),
    "page_questions": table("02_Page_Questions", 4, 5, 20, 10),
    "content_id": table("03_Content_ID", 4, 5, 48, 17),
    "content_en": table("04_Content_EN", 4, 5, 48, 17),
    "proof_placeholders": table("05_Proof_Placeholders", 4, 5, 25, 10),
    "positioning_qa": table("06_Positioning_Tone_QA", 4, 5, 17, 12),
    "approval_actions": table("07_Approval_Actions", 4, 5, 20, 9),
    "source_notes": table("08_Source_Notes", 4, 5, 16, 5),
}

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
OUTPUT.write_text(json.dumps(data, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
print(json.dumps({key: len(value) for key, value in data.items()}, ensure_ascii=False, indent=2))
