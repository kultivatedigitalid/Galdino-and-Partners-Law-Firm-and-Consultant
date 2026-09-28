import json
import zipfile
from pathlib import Path
from openpyxl import load_workbook

path = Path(r"C:\Users\Joshua\OneDrive\Documents\Law\outputs\galdino_partner\fase2\GP_Fase2_Website_Content_Sheet_Mapping_Aligned_2026-08-14.xlsx")
expected = [
    "00_Dashboard", "01_GLOBAL", "02_Page_Questions", "03_Content_ID", "04_Content_EN",
    "05_Proof_Placeholders", "06_Positioning_Tone_QA", "07_Approval_Actions", "08_Source_Notes",
    "09_Page_Section_Mapping", "10_Asset_Register", "11_Reference_Register", "12_Mapping_QA",
]
with zipfile.ZipFile(path, "r") as archive:
    bad_zip_member = archive.testzip()

workbook = load_workbook(path, read_only=False, data_only=False)
mapping = workbook["09_Page_Section_Mapping"]
assets = workbook["10_Asset_Register"]
references = workbook["11_Reference_Register"]
qa = workbook["12_Mapping_QA"]

def nonempty_ids(sheet, start_row=5):
    return [sheet.cell(row, 1).value for row in range(start_row, sheet.max_row + 1) if sheet.cell(row, 1).value not in (None, "")]

mapping_ids = nonempty_ids(mapping)
asset_ids = nonempty_ids(assets)
reference_ids = nonempty_ids(references)
qa_ids = [qa.cell(row, 1).value for row in range(5, 18) if qa.cell(row, 1).value]
formula_count = sum(1 for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row if isinstance(cell.value, str) and cell.value.startswith("="))

result = {
    "path": str(path), "sizeBytes": path.stat().st_size,
    "zipIntegrity": "PASS" if bad_zip_member is None else f"FAIL: {bad_zip_member}",
    "sheetNamesMatch": workbook.sheetnames == expected, "sheetCount": len(workbook.sheetnames),
    "mappingRows": len(mapping_ids), "assetRows": len(asset_ids), "referenceRows": len(reference_ids), "qaRows": len(qa_ids),
    "uniqueMappingIDs": len(mapping_ids) == len(set(mapping_ids)), "uniqueAssetIDs": len(asset_ids) == len(set(asset_ids)),
    "uniqueReferenceIDs": len(reference_ids) == len(set(reference_ids)), "uniqueQAIDs": len(qa_ids) == len(set(qa_ids)),
    "serviceDetailsFormula": qa["G14"].value, "finalMappedFormula": qa["B25"].value,
    "formulaCount": formula_count,
    "row3Heights": {sheet.title: sheet.row_dimensions[3].height for sheet in [mapping, assets, references, qa]},
    "mergedNotesRemoved": all(not any(str(rng).startswith("A2:") or str(rng).startswith("A3:") for rng in sheet.merged_cells.ranges) for sheet in [mapping, assets, references, qa]),
}
print(json.dumps(result, ensure_ascii=False, indent=2))
checks = [
    bad_zip_member is None, workbook.sheetnames == expected, len(mapping_ids) == 46, len(asset_ids) == 18,
    len(reference_ids) == 27, len(qa_ids) == 13, result["uniqueMappingIDs"], result["uniqueAssetIDs"],
    result["uniqueReferenceIDs"], result["uniqueQAIDs"], result["mergedNotesRemoved"],
]
raise SystemExit(0 if all(checks) else 1)
