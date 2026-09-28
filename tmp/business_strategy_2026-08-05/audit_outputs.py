from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader


ROOT = Path(r"C:\Users\Joshua\OneDrive\Documents\Law")
OUT = ROOT / "outputs" / "galdino_partner" / "business_strategy_2026-08-05"
DOCX = OUT / "GP_Business_Service_Portfolio_Strategy_2026-08-05.docx"
XLSX = OUT / "GP_Service_Portfolio_Scenario_Matrix_2026-08-05.xlsx"
PDF = ROOT / "tmp" / "business_strategy_2026-08-05" / "docx_render" / "GP_Business_Service_Portfolio_Strategy_2026-08-05.pdf"


def zip_audit(path: Path) -> dict:
    with zipfile.ZipFile(path) as archive:
        bad = archive.testzip()
        names = archive.namelist()
    return {"valid_zip": bad is None, "bad_member": bad, "members": len(names)}


def docx_audit() -> dict:
    doc = Document(DOCX)
    text = "\n".join(p.text for p in doc.paragraphs)
    section = doc.sections[0]
    styles = {}
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        styles[name] = {
            "font": style.font.name,
            "size_pt": style.font.size.pt if style.font.size else None,
        }
    marker_patterns = [r"\{\{.+?\}\}", r"\bTBD\b", r"\bTODO\b", r"lorem ipsum"]
    markers = []
    for pattern in marker_patterns:
        markers.extend(re.findall(pattern, text, flags=re.I))
    empty_tables = [i + 1 for i, table in enumerate(doc.tables) if not any(cell.text.strip() for row in table.rows for cell in row.cells)]
    return {
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "empty_tables": empty_tables,
        "sections": len(doc.sections),
        "page_inches": [round(section.page_width.inches, 3), round(section.page_height.inches, 3)],
        "margins_inches": {
            "top": round(section.top_margin.inches, 3),
            "right": round(section.right_margin.inches, 3),
            "bottom": round(section.bottom_margin.inches, 3),
            "left": round(section.left_margin.inches, 3),
        },
        "styles": styles,
        "unresolved_markers": markers,
        "zip": zip_audit(DOCX),
    }


def xlsx_audit() -> dict:
    wb_formula = load_workbook(XLSX, data_only=False, read_only=False)
    wb_values = load_workbook(XLSX, data_only=True, read_only=False)
    formula_count = 0
    formula_error_values = []
    sheets = []
    expected_errors = {"#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A"}
    for ws in wb_formula.worksheets:
        value_ws = wb_values[ws.title]
        formulas = 0
        errors = []
        nonempty = 0
        for row in ws.iter_rows():
            for cell in row:
                if cell.value not in (None, ""):
                    nonempty += 1
                if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")):
                    formulas += 1
                    formula_count += 1
                    cached = value_ws[cell.coordinate].value
                    if cached in expected_errors:
                        error = f"{ws.title}!{cell.coordinate}:{cached}"
                        errors.append(error)
                        formula_error_values.append(error)
        sheets.append({"name": ws.title, "rows": ws.max_row, "columns": ws.max_column, "nonempty": nonempty, "formulas": formulas, "formula_errors": errors})
    return {
        "sheet_count": len(wb_formula.sheetnames),
        "sheet_names": wb_formula.sheetnames,
        "formula_count": formula_count,
        "formula_error_values": formula_error_values,
        "empty_sheets": [s["name"] for s in sheets if s["nonempty"] == 0],
        "sheets": sheets,
        "zip": zip_audit(XLSX),
    }


def pdf_audit() -> dict:
    reader = PdfReader(PDF)
    chars_per_page = [len((page.extract_text() or "").strip()) for page in reader.pages]
    return {
        "pages": len(reader.pages),
        "characters_per_page": chars_per_page,
        "blank_pages": [i + 1 for i, count in enumerate(chars_per_page) if count < 25],
    }


result = {
    "docx": docx_audit(),
    "xlsx": xlsx_audit(),
    "pdf_render": pdf_audit(),
}
print(json.dumps(result, ensure_ascii=False, indent=2))
